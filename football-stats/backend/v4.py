import undetected_chromedriver as uc
from bs4 import BeautifulSoup
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import time

# =====================================================
# JOGADORES — separados por tipo
# =====================================================

goleiros = [
    {"name": "Mike Maignan",  "url": "https://www.transfermarkt.com.br/mike-maignan/leistungsdatendetails/spieler/182906/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Brice Samba",   "url": "https://www.transfermarkt.com.br/brice-samba/leistungsdatendetails/spieler/191056/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Robin Risser",  "url": "https://www.transfermarkt.com.br/robin-risser/leistungsdatendetails/spieler/743515/saison/2025/wettbewerb/0/plus/1"},
]

linha = [
    {"name": "Lucas Digne",         "url": "https://www.transfermarkt.com.br/lucas-digne/leistungsdatendetails/spieler/126664/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Jules Kounde",        "url": "https://www.transfermarkt.com.br/jules-kounde/leistungsdatendetails/spieler/412517/saison/2025/wettbewerb/0/plus/1"},
    {"name": "William Saliba",      "url": "https://www.transfermarkt.com.br/william-saliba/leistungsdatendetails/spieler/495666/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Ibrahima Konate",     "url": "https://www.transfermarkt.com.br/ibrahima-konate/leistungsdatendetails/spieler/357119/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Dayot Upamecano",     "url": "https://www.transfermarkt.com.br/dayot-upamecano/leistungsdatendetails/spieler/344695/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Aurelien Tchouameni", "url": "https://www.transfermarkt.com.br/aurelien-tchouameni/leistungsdatendetails/spieler/413112/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Eduardo Camavinga",   "url": "https://www.transfermarkt.com.br/eduardo-camavinga/leistungsdatendetails/spieler/640428/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Warren Zaire Emery",  "url": "https://www.transfermarkt.com.br/warren-zaire-emery/leistungsdatendetails/spieler/810092/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Kylian Mbappe",       "url": "https://www.transfermarkt.com.br/kylian-mbappe/leistungsdatendetails/spieler/342229/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Ousmane Dembele",     "url": "https://www.transfermarkt.com.br/ousmane-dembele/leistungsdatendetails/spieler/288230/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Marcus Thuram",       "url": "https://www.transfermarkt.com.br/marcus-thuram/leistungsdatendetails/spieler/318528/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Randal Kolo Muani",   "url": "https://www.transfermarkt.com.br/randal-kolo-muani/leistungsdatendetails/spieler/487969/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Bradley Barcola",     "url": "https://www.transfermarkt.com.br/bradley-barcola/leistungsdatendetails/spieler/708265/saison/2025/wettbewerb/0/plus/1"},
    {"name": "Kingsley Coman",      "url": "https://www.transfermarkt.com.br/kingsley-coman/leistungsdatendetails/spieler/243714/saison/2025/wettbewerb/0/plus/1"},
]

# =====================================================
# SCRAPING — helpers
# =====================================================

def cell_text(cell):
    t = cell.get_text(separator=" ", strip=True)
    return t if t else "-"

def scrape_player(driver, player):
    """Abre a página, scrolla e retorna o soup."""
    driver.get(player["url"])
    time.sleep(5)
    for _ in range(6):
        driver.execute_script("window.scrollBy(0, 900);")
        time.sleep(1.5)
    driver.execute_script("window.scrollTo(0, document.body.scrollHeight / 2);")
    time.sleep(3)
    return BeautifulSoup(driver.page_source, "html.parser")

def parse_resumo_goleiro(row, name):
    """
    Tabela de média — GOLEIRO
    Colunas TM (div[role='row']):
      [0] Competição
      [1] Jogos Relacionados
      [2] Partidas Jogadas (link — ignorado)
      [3] PPJ
      [4] Gols (sempre -)
      [5] Gols Contra (sempre -)
      [6] Entrou
      [7] Saiu
      [8] Cartão Amarelo
      [9] Expulso 2 Amarelos
      [10] Cartão Vermelho
      [11] Gols Sofridos
      [12] Jogos sem Gols Sofridos
      [13] Minutos Jogados
    """
    cells = row.select(".tm-grid__cell")
    t = [cell_text(c) for c in cells]
    if len(t) < 10:
        return None
    # só linhas de resumo: PPJ não tem ":"
    if ":" in str(t[3]):
        return None
    return {
        "Jogador":              name,
        "Competicao":           t[0],
        "Jogos Relacionados":   t[1],
        "Partidas Jogadas":     t[1],   # mesmo valor, link é ignorado
        "PPJ":                  t[3],
        "Gols":                 t[4]  if len(t) > 4  else "-",
        "Gols Contra":          t[5]  if len(t) > 5  else "-",
        "Entrou":               t[6]  if len(t) > 6  else "-",
        "Saiu":                 t[7]  if len(t) > 7  else "-",
        "Cartao Amarelo":       t[8]  if len(t) > 8  else "-",
        "Expulso 2 Amarelos":   t[9]  if len(t) > 9  else "-",
        "Cartao Vermelho":      t[10] if len(t) > 10 else "-",
        "Gols Sofridos":        t[11] if len(t) > 11 else "-",
        "Sem Gols Sofridos":    t[12] if len(t) > 12 else "-",
        "Minutos Jogados":      t[13] if len(t) > 13 else t[-1],
    }

def parse_jogo_goleiro(row, name):
    """
    Tabela de jogos — GOLEIRO
    Mesmo div[role='row'] mas PPJ tem ":"
    Mapeamento real das células para jogos individuais:
      [0] Rodada
      [1] Data
      [2] Cidade
      [3] PPJ (placar ex: 1:2)  → Resultado
      [4] Posição
      [8] Cartão Amarelo
      [9] Expulso 2 Amarelos
      [10] Cartão Vermelho
      [12] Entrou (min)
      [13] Minutos Jogados
    (Time casa/visitante não vêm nas células de texto do div, ficam como -)
    """
    cells = row.select(".tm-grid__cell")
    t = [cell_text(c) for c in cells]
    if len(t) < 10:
        return None
    if ":" not in str(t[3]):
        return None
    return {
        "Jogador":            name,
        "Rodada":             t[0],
        "Data":               t[1],
        "Cidade":             t[2],
        "Time Casa":          "-",
        "Time Visitante":     "-",
        "Resultado":          t[3],
        "Posicao":            t[4]  if len(t) > 4  else "-",
        "Gols":               t[5]  if len(t) > 5  else "-",
        "Assistencias":       t[6]  if len(t) > 6  else "-",
        "Gols Contra":        t[7]  if len(t) > 7  else "-",
        "Cartao Amarelo":     t[8]  if len(t) > 8  else "-",
        "Expulso 2 Amarelos": t[9]  if len(t) > 9  else "-",
        "Cartao Vermelho":    t[10] if len(t) > 10 else "-",
        "Entrou":             t[12] if len(t) > 12 else "-",
        "Saiu":               "-",
        "Minutos":            t[13] if len(t) > 13 else t[-1],
    }

def parse_resumo_linha(row, name):
    """
    Tabela de média — JOGADOR DE LINHA
    Colunas TM:
      [0] Competição
      [1] Jogos Relacionados
      [2] link (ignorado)
      [3] PPJ
      [4] Gols
      [5] Assistências
      [6] Gols Contra
      [7] Cartão Amarelo
      [8] Expulso 2 Amarelos
      [9] Cartão Vermelho
      [10] Entrou
      [11] Saiu
      [12] Gols de Pênalti
      [13] Minutos por Gol
      [14] Minutos Jogados
    """
    cells = row.select(".tm-grid__cell")
    t = [cell_text(c) for c in cells]
    if len(t) < 10:
        return None
    if ":" in str(t[3]):
        return None
    return {
        "Jogador":            name,
        "Competicao":         t[0],
        "Jogos Relacionados": t[1],
        "PPJ":                t[3],
        "Gols":               t[4]  if len(t) > 4  else "-",
        "Assistencias":       t[5]  if len(t) > 5  else "-",
        "Gols Contra":        t[6]  if len(t) > 6  else "-",
        "Cartao Amarelo":     t[7]  if len(t) > 7  else "-",
        "Expulso 2 Amarelos": t[8]  if len(t) > 8  else "-",
        "Cartao Vermelho":    t[9]  if len(t) > 9  else "-",
        "Entrou":             t[10] if len(t) > 10 else "-",
        "Saiu":               t[11] if len(t) > 11 else "-",
        "Gols Penalti":       t[12] if len(t) > 12 else "-",
        "Minutos por Gol":    t[13] if len(t) > 13 else "-",
        "Minutos Jogados":    t[14] if len(t) > 14 else t[-1],
    }

def parse_jogo_linha(row, name):
    """
    Tabela de jogos — JOGADOR DE LINHA
    Mapeamento real das células para jogos individuais:
      [0] Rodada  [1] Data  [2] Cidade  [3] Resultado (placar)
      [4] Posição  [5] Gols  [6] Assistências  [7] Gols Contra
      [8] Cartão Amarelo  [9] Expulso 2 Amarelos  [10] Cartão Vermelho
      [11] Entrou  [12] Saiu  [-1] Minutos
    """
    cells = row.select(".tm-grid__cell")
    t = [cell_text(c) for c in cells]
    if len(t) < 10:
        return None
    if ":" not in str(t[3]):
        return None
    return {
        "Jogador":            name,
        "Rodada":             t[0],
        "Data":               t[1],
        "Cidade":             t[2],
        "Time Casa":          "-",
        "Time Visitante":     "-",
        "Resultado":          t[3],
        "Posicao":            t[4]  if len(t) > 4  else "-",
        "Gols":               t[5]  if len(t) > 5  else "-",
        "Assistencias":       t[6]  if len(t) > 6  else "-",
        "Gols Contra":        t[7]  if len(t) > 7  else "-",
        "Cartao Amarelo":     t[8]  if len(t) > 8  else "-",
        "Expulso 2 Amarelos": t[9]  if len(t) > 9  else "-",
        "Cartao Vermelho":    t[10] if len(t) > 10 else "-",
        "Entrou":             t[11] if len(t) > 11 else "-",
        "Saiu":               t[12] if len(t) > 12 else "-",
        "Minutos":            t[-1],
    }

# =====================================================
# MAIN LOOP
# =====================================================

options = uc.ChromeOptions()
options.add_argument("--start-maximized")
driver = uc.Chrome(options=options, version_main=148)

resumo_gol, jogos_gol   = [], []
resumo_lin, jogos_lin   = [], []

for player in goleiros:
    print(f"\n{'='*45}\n  [GOL] {player['name']}\n{'='*45}")
    try:
        soup = scrape_player(driver, player)
        for row in soup.select('div[role="row"][aria-rowindex]'):
            r = parse_resumo_goleiro(row, player["name"])
            j = parse_jogo_goleiro(row, player["name"])
            if r: resumo_gol.append(r)
            if j: jogos_gol.append(j)
        print(f"  resumo: {sum(1 for x in resumo_gol if x['Jogador']==player['name'])} | jogos: {sum(1 for x in jogos_gol if x['Jogador']==player['name'])}")
    except Exception as e:
        print(f"  ERRO: {e}")

for player in linha:
    print(f"\n{'='*45}\n  [LIN] {player['name']}\n{'='*45}")
    try:
        soup = scrape_player(driver, player)
        for row in soup.select('div[role="row"][aria-rowindex]'):
            r = parse_resumo_linha(row, player["name"])
            j = parse_jogo_linha(row, player["name"])
            if r: resumo_lin.append(r)
            if j: jogos_lin.append(j)
        print(f"  resumo: {sum(1 for x in resumo_lin if x['Jogador']==player['name'])} | jogos: {sum(1 for x in jogos_lin if x['Jogador']==player['name'])}")
    except Exception as e:
        print(f"  ERRO: {e}")

driver.quit()

# =====================================================
# DATAFRAMES
# =====================================================

def make_df(rows, cols):
    df = pd.DataFrame(rows) if rows else pd.DataFrame(columns=cols)
    for c in cols:
        if c not in df.columns:
            df[c] = "-"
    return df[cols].reset_index(drop=True)

COLS_RESUMO_GOL = [
    "Jogador", "Competicao", "Jogos Relacionados", "PPJ",
    "Gols", "Gols Contra",
    "Entrou", "Saiu",
    "Cartao Amarelo", "Expulso 2 Amarelos", "Cartao Vermelho",
    "Gols Sofridos", "Sem Gols Sofridos",
    "Minutos Jogados",
]

COLS_JOGOS_GOL = [
    "Jogador", "Rodada", "Data", "Cidade",
    "Time Casa", "Time Visitante", "Resultado",
    "Posicao", "Gols", "Assistencias", "Gols Contra",
    "Cartao Amarelo", "Expulso 2 Amarelos", "Cartao Vermelho",
    "Entrou", "Saiu", "Minutos",
]

COLS_RESUMO_LIN = [
    "Jogador", "Competicao", "Jogos Relacionados", "PPJ",
    "Gols", "Assistencias", "Gols Contra",
    "Cartao Amarelo", "Expulso 2 Amarelos", "Cartao Vermelho",
    "Entrou", "Saiu",
    "Gols Penalti", "Minutos por Gol", "Minutos Jogados",
]

COLS_JOGOS_LIN = [
    "Jogador", "Rodada", "Data", "Cidade",
    "Time Casa", "Time Visitante", "Resultado",
    "Posicao", "Gols", "Assistencias", "Gols Contra",
    "Cartao Amarelo", "Expulso 2 Amarelos", "Cartao Vermelho",
    "Entrou", "Saiu", "Minutos",
]

df_rg = make_df(resumo_gol, COLS_RESUMO_GOL)
df_jg = make_df(jogos_gol,  COLS_JOGOS_GOL)
df_rl = make_df(resumo_lin, COLS_RESUMO_LIN)
df_jl = make_df(jogos_lin,  COLS_JOGOS_LIN)

# limpa resumos
for df in [df_rg, df_rl]:
    df.drop_duplicates(subset=["Jogador", "Competicao"], inplace=True)
    df.reset_index(drop=True, inplace=True)

print(f"\nResumo Gol: {len(df_rg)} | Jogos Gol: {len(df_jg)}")
print(f"Resumo Lin: {len(df_rl)} | Jogos Lin: {len(df_jl)}")

# =====================================================
# EXCEL — helpers
# =====================================================

NAVY = "1B3251"; WHITE = "FFFFFF"; LIGHT_BG = "F0F4FA"; LIGHT_GOL = "FFF8E7"

def hdr_font(size=8, color=WHITE):
    return Font(name="Arial", bold=True, color=color, size=size)

def body_font(bold=False):
    return Font(name="Arial", bold=bold, color="1A1A1A", size=9)

def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def center(wrap=False):
    return Alignment(horizontal="center", vertical="center", wrap_text=wrap)

def left():
    return Alignment(horizontal="left", vertical="center")

thin  = Side(style="thin",   color="C0CCE0")
thick = Side(style="medium", color=NAVY)
b_all = Border(left=thin, right=thin, top=thin, bottom=thin)
b_hdr = Border(left=thin, right=thin, top=thin, bottom=thick)

def write_header(ws, grupos, headers, offset):
    for grp in grupos:
        cs = grp["cols"][0] + offset
        ce = grp["cols"][-1] + offset
        ws.merge_cells(start_row=1, start_column=cs, end_row=1, end_column=ce)
        c = ws.cell(row=1, column=cs, value=grp["label"])
        c.font = hdr_font(); c.fill = fill(grp["cor"])
        c.alignment = center(); c.border = b_all
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=2, column=i + offset, value=h)
        c.font = hdr_font(); c.fill = fill(NAVY)
        c.alignment = center(wrap=True); c.border = b_hdr
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 32

def write_rows(ws, df, cols, offset, bg_alt, result_col=None):
    jogador_atual = None; flip = False
    for row_i, (_, row) in enumerate(df.iterrows(), start=3):
        if row["Jogador"] != jogador_atual:
            jogador_atual = row["Jogador"]; flip = not flip
        bg = bg_alt if flip else WHITE
        for ci, cn in enumerate(cols, 1):
            val = row.get(cn, "-")
            val = "-" if (val != val or str(val).strip() == "") else val
            c = ws.cell(row=row_i, column=ci + offset, value=val)
            c.fill = fill(bg); c.border = b_all
            c.font = body_font(bold=(ci == 1))
            if cn == result_col:
                s = str(val)
                if ":" in s:
                    try:
                        a, b = int(s.split(":")[0]), int(s.split(":")[1])
                        cor = "145A32" if a > b else ("7B241C" if a < b else "784212")
                        c.font = Font(name="Arial", bold=True, color=cor, size=9)
                    except: pass
                c.alignment = center()
            else:
                c.alignment = left() if ci <= 2 else center()
        ws.row_dimensions[row_i].height = 15

def set_widths(ws, widths, offset):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i + offset)].width = w

def set_separator(ws, col):
    ws.column_dimensions[get_column_letter(col)].width = 3

# =====================================================
# EXCEL — ABA GOLEIROS
# =====================================================

wb = openpyxl.Workbook()
ws_gol = wb.active
ws_gol.title = "Goleiros"
ws_gol.freeze_panes = "A3"

# — Tabela 1: Resumo Goleiros (offset 0) —
N_RG = len(COLS_RESUMO_GOL)   # 14 colunas

grp_rg = [
    {"cols": [1,2,3,4],   "cor": NAVY,    "label": "IDENTIFICAÇÃO"},
    {"cols": [5,6],       "cor": "1A5276", "label": "ATAQUE"},
    {"cols": [7,8],       "cor": "145A32", "label": "PARTICIPAÇÃO"},
    {"cols": [9,10,11],   "cor": "7B241C", "label": "CARTÕES"},
    {"cols": [12,13],     "cor": "784212", "label": "GOLEIRO"},
    {"cols": [14],        "cor": "2C3E50", "label": "TEMPO"},
]
hdr_rg = [
    "Jogador", "Competição", "Jogos\nRelac.", "PPJ",
    "Gols", "Gols\nContra",
    "Entrou", "Saiu",
    "Cartão\nAmarelo", "Expulso\n2 Amar.", "Cartão\nVermelho",
    "Gols\nSofridos", "Sem Gols\nSofridos",
    "Minutos\nJogados",
]

write_header(ws_gol, grp_rg, hdr_rg, offset=0)
write_rows(ws_gol, df_rg, COLS_RESUMO_GOL, offset=0, bg_alt=LIGHT_GOL)
set_widths(ws_gol, [20,22,9,7, 7,9, 8,7, 10,10,10, 10,13, 10], offset=0)
set_separator(ws_gol, N_RG + 1)

# — Tabela 2: Jogos Goleiros (offset = N_RG + 1) —
OFF_JG = N_RG + 1

grp_jg = [
    {"cols": [1,2,3,4,5,6,7], "cor": NAVY,    "label": "IDENTIFICAÇÃO DO JOGO"},
    {"cols": [8],              "cor": "145A32", "label": "POS."},
    {"cols": [9,10,11],        "cor": "1A5276", "label": "CONTRIBUIÇÃO"},
    {"cols": [12,13,14],       "cor": "7B241C", "label": "CARTÕES"},
    {"cols": [15,16,17],       "cor": "2C3E50", "label": "TEMPO"},
]
hdr_jg = [
    "Jogador", "Rodada", "Data", "Cidade", "Time Casa", "T. Visitante", "Resultado",
    "Posição",
    "Gols", "Assist.", "Gols\nContra",
    "Amarelo", "Exp.\n2 Amar.", "Vermelho",
    "Entrou\n(min)", "Saiu\n(min)", "Minutos",
]

write_header(ws_gol, grp_jg, hdr_jg, offset=OFF_JG)
write_rows(ws_gol, df_jg, COLS_JOGOS_GOL, offset=OFF_JG, bg_alt=LIGHT_GOL, result_col="Resultado")
set_widths(ws_gol, [20,8,11,10,14,14,10, 9, 7,8,9, 8,9,8, 9,9,10], offset=OFF_JG)

# =====================================================
# EXCEL — ABA JOGADORES DE LINHA
# =====================================================

ws_lin = wb.create_sheet("Jogadores de Linha")
ws_lin.freeze_panes = "A3"

N_RL = len(COLS_RESUMO_LIN)   # 15 colunas

# — Tabela 3: Resumo Linha (offset 0) —
grp_rl = [
    {"cols": [1,2,3,4],    "cor": NAVY,    "label": "IDENTIFICAÇÃO"},
    {"cols": [5,6,7],      "cor": "1A5276", "label": "ATAQUE"},
    {"cols": [8,9,10],     "cor": "7B241C", "label": "CARTÕES"},
    {"cols": [11,12],      "cor": "145A32", "label": "PARTICIPAÇÃO"},
    {"cols": [13,14,15],   "cor": "2C3E50", "label": "TEMPO / BÔNUS"},
]
hdr_rl = [
    "Jogador", "Competição", "Jogos\nRelac.", "PPJ",
    "Gols", "Assist.", "Gols\nContra",
    "Cartão\nAmarelo", "Expulso\n2 Amar.", "Cartão\nVermelho",
    "Entrou", "Saiu",
    "Gols\nPênalti", "Min. p/\nGol", "Minutos\nJogados",
]

write_header(ws_lin, grp_rl, hdr_rl, offset=0)
write_rows(ws_lin, df_rl, COLS_RESUMO_LIN, offset=0, bg_alt=LIGHT_BG)
set_widths(ws_lin, [20,22,9,7, 7,8,9, 10,10,10, 8,7, 9,9,10], offset=0)
set_separator(ws_lin, N_RL + 1)

# — Tabela 4: Jogos Linha (offset = N_RL + 1) —
OFF_JL = N_RL + 1

grp_jl = [
    {"cols": [1,2,3,4,5,6,7], "cor": NAVY,    "label": "IDENTIFICAÇÃO DO JOGO"},
    {"cols": [8],              "cor": "145A32", "label": "POS."},
    {"cols": [9,10,11],        "cor": "1A5276", "label": "CONTRIBUIÇÃO"},
    {"cols": [12,13,14],       "cor": "7B241C", "label": "CARTÕES"},
    {"cols": [15,16,17],       "cor": "2C3E50", "label": "TEMPO"},
]
hdr_jl = [
    "Jogador", "Rodada", "Data", "Cidade", "Time Casa", "T. Visitante", "Resultado",
    "Posição",
    "Gols", "Assist.", "Gols\nContra",
    "Amarelo", "Exp.\n2 Amar.", "Vermelho",
    "Entrou\n(min)", "Saiu\n(min)", "Minutos",
]

write_header(ws_lin, grp_jl, hdr_jl, offset=OFF_JL)
write_rows(ws_lin, df_jl, COLS_JOGOS_LIN, offset=OFF_JL, bg_alt=LIGHT_BG, result_col="Resultado")
set_widths(ws_lin, [20,8,11,10,14,14,10, 9, 7,8,9, 8,9,8, 9,9,10], offset=OFF_JL)

# =====================================================
# SALVAR
# =====================================================

wb.save("football_stats_limpo.xlsx")
print("\n============================================")
print("EXCEL SALVO: football_stats_limpo.xlsx")
print("Aba 'Goleiros':          Resumo | Jogos")
print("Aba 'Jogadores de Linha': Resumo | Jogos")
print("============================================")